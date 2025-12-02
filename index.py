import json
import os
import unicodedata
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment
import tkinter as tk
from tkinter import ttk, messagebox, simpledialog

ARCHIVO = "db.json"

PRECIO_CAJA = 1000  # 🔹 Precio inicial de la bandeja de huevos
DEFAULT_CAJA_MANUAL = {}

# ------------------ Funciones base ------------------

def normalizar(texto):
    return ''.join(
        c for c in unicodedata.normalize('NFD', (texto or "").lower())
        if unicodedata.category(c) != 'Mn'
    )

def cargar_datos():
    if not os.path.exists(ARCHIVO):
        return {
            "clientes": [],
            "precio_caja": PRECIO_CAJA,
            "precios_por_comuna": {},
            "movimientos": [],
            "caja_manual": DEFAULT_CAJA_MANUAL.copy(),
            "comunas": []
        }
    try:
        with open(ARCHIVO, "r", encoding="utf-8") as f:
            data = json.load(f)
            if not isinstance(data, dict):
                raise ValueError("El archivo JSON no tiene la estructura esperada.")
            # Asegurarse de que las claves necesarias estén presentes
            data.setdefault("clientes", [])
            data.setdefault("precio_caja", PRECIO_CAJA)
            data.setdefault("precios_por_comuna", {})
            data.setdefault("movimientos", [])
            data.setdefault("comunas", [])
            data["caja_manual"] = {}
            return data
    except (json.JSONDecodeError, ValueError, IOError):
        messagebox.showerror("Error", "El archivo de datos está corrupto o tiene un formato incorrecto. Se restablecerán los valores predeterminados.")
        return {
            "clientes": [],
            "precio_caja": PRECIO_CAJA,
            "precios_por_comuna": {},
            "movimientos": [],
            "caja_manual": DEFAULT_CAJA_MANUAL.copy(),
            "comunas": []
        }

def guardar_datos(data):
    # Cargar los datos existentes para preservar las claves
    datos_existentes = cargar_datos()
    datos_existentes.update(data)  # Actualizar solo las claves proporcionadas
    with open(ARCHIVO, "w", encoding="utf-8") as f:
        json.dump(datos_existentes, f, indent=4, ensure_ascii=False)

# ------------------ Interfaz gráfica ------------------

class App:
    def __init__(self, root):
        self.root = root
        self.style = ttk.Style()
        self.style.theme_use("clam")

        self.temas = {
            "claro": {
                "bg": "#f0f4f8",
                "panel": "#f7f9fb",
                "texto": "#003366",
                "texto_secundario": "#4b5d6b",
                "tree_bg": "#ffffff",
                "tree_fg": "#111827",
                "tree_sel": "#cce5ff",
                "tree_header_bg": "#d1e7ff",
                "tree_header_fg": "#003366",
                "entry_bg": "#ffffff",
                "entry_fg": "#111827",
                "placeholder": "#777777",
                "boton_bg": "#e2ecff",
                "boton_fg": "#003366",
                "boton_hover": "#d0dcf5"
            },
            "oscuro": {
                "bg": "#1f2430",
                "panel": "#2a2f3d",
                "texto": "#e6edf7",
                "texto_secundario": "#9da7c5",
                "tree_bg": "#272c38",
                "tree_fg": "#f7f9fe",
                "tree_sel": "#3d4558",
                "tree_header_bg": "#303644",
                "tree_header_fg": "#e6edf7",
                "entry_bg": "#32394a",
                "entry_fg": "#f7f9fe",
                "placeholder": "#8c94ab",
                "boton_bg": "#3a4255",
                "boton_fg": "#f0f4ff",
                "boton_hover": "#4a546b"
            }
        }
        self.tema_actual = "claro"
        self.colores = self.temas[self.tema_actual]
        self.widgets_tema = []
        self.root.configure(bg=self.colores["bg"])
        self.registrar_widget_tema(self.root, fondo="bg")
        self.configurar_estilos_ttk()

        self.root.title("📦 Control de Reparto de Huevos")
        self.root.geometry("1024x720")
        self.root.configure(bg=self.colores["bg"])

        datos_cargados = cargar_datos()
        self.data = datos_cargados.get("clientes", [])
        self.precios_por_comuna = {
            self.estandarizar_comuna(comuna): precio
            for comuna, precio in datos_cargados.get("precios_por_comuna", {}).items()
            if self.estandarizar_comuna(comuna)
        }
        self.movimientos = datos_cargados.get("movimientos", [])
        self.caja_manual = dict(datos_cargados.get("caja_manual", {}))
        self._combobox_comunas = []
        self._comunas_map = {}
        self.comunas = []
        self.filtro_comuna_actual = None
        self.filtro_dia_actual = None
        self.actualizar_comunas_existentes(datos_cargados.get("comunas", []))
        global PRECIO_CAJA
        PRECIO_CAJA = datos_cargados.get("precio_caja", PRECIO_CAJA)

        frame = self.crear_frame_tema(root, fondo="bg", padx=16, pady=16)
        frame.pack(fill="both", expand=True)

        top_bar = self.crear_frame_tema(frame, fondo="bg")
        top_bar.pack(fill="x")

        self.tema_var = tk.BooleanVar(value=self.tema_actual == "oscuro")
        self.boton_tema = ttk.Checkbutton(
            top_bar,
            text="☀️",
            style="Switch.TCheckbutton",
            width=4,
            variable=self.tema_var,
            command=self.alternar_tema
        )
        self.boton_tema.pack(side="right", padx=(0, 4))
        self.boton_tema.config(compound="center")
        self.actualizar_texto_boton_tema()

        titulo = self.crear_label_tema(frame, "📋 Control de Reparto", font=("Segoe UI", 18, "bold"))
        titulo.pack(pady=(4, 12))

        botones_frame = self.crear_frame_tema(frame, fondo="bg")
        botones_frame.pack(pady=12)

        botones = [
            ("Agregar cliente", self.ventana_agregar_cliente),
            ("Nuevo pedido", self.ventana_nuevo_pedido),
            ("Editar", self.ventana_editar),
            ("Generar reparto", self.generar_reparto),
            ("Ver resumen", self.ventana_resumen),
            ("Gestionar Precios", self.gestionar_precios_por_comuna),
            ("Cambiar precio", self.cambiar_precio_caja),
            ("Gestión de Caja", self.ventana_caja),
            ("Salir", root.quit)
        ]

        for text, cmd in botones:
            ttk.Button(botones_frame, text=text, command=cmd, width=20).pack(side="left", padx=6)

        filtros_frame = self.crear_frame_tema(frame, fondo="bg")
        filtros_frame.pack(fill="x", pady=(4, 10))

        comuna_container = self.crear_frame_tema(filtros_frame, fondo="bg")
        comuna_container.pack(side="left")
        self.crear_label_tema(comuna_container, "Comuna", font=("Segoe UI", 9)).pack(side="left", padx=(0, 6))
        self.combo_filtro_comuna = self.crear_combobox_comunas(
            comuna_container,
            permitir_agregar=False,
            incluir_todas=True,
            width=22
        )
        self.combo_filtro_comuna.pack(side="left")

        dia_container = self.crear_frame_tema(filtros_frame, fondo="bg")
        dia_container.pack(side="left", padx=(18, 0))
        self.crear_label_tema(dia_container, "Día", font=("Segoe UI", 9)).pack(side="left", padx=(0, 6))
        self.combo_filtro_dia = ttk.Combobox(dia_container, state="readonly", width=18)
        self.combo_filtro_dia.pack(side="left")

        limpiar_container = self.crear_frame_tema(filtros_frame, fondo="bg")
        limpiar_container.pack(side="left", padx=(18, 0))
        ttk.Button(
            limpiar_container,
            text="↺ Limpiar",
            command=self.restablecer_filtros,
            width=10
        ).pack(side="left")

        def on_cambio_comuna(_):
            seleccion = self.combo_filtro_comuna.get()
            if seleccion == "Todas" or not seleccion:
                self.filtro_comuna_actual = None
            else:
                self.filtro_comuna_actual = self.registrar_comuna(seleccion, actualizar_opciones=False)
            self.ver_clientes()

        def on_cambio_dia(_):
            seleccion = self.combo_filtro_dia.get()
            if seleccion == "Todos" or not seleccion:
                self.filtro_dia_actual = None
            else:
                self.filtro_dia_actual = seleccion
            self.ver_clientes()

        self.combo_filtro_comuna.bind("<<ComboboxSelected>>", on_cambio_comuna, add="+")
        self.combo_filtro_dia.bind("<<ComboboxSelected>>", on_cambio_dia, add="+")

        self.actualizar_opciones_dias()
        self.restablecer_filtros(actualizar_tabla=False)

        self.tree = ttk.Treeview(
            frame,
            columns=("Nombre", "Teléfono", "Dirección", "Comuna", "Día de Reparto", "Pendiente a entrega", "Total histórico", "Total adeudado"),
            show="headings",
            height=16
        )

        for col in ("Nombre", "Teléfono", "Dirección", "Comuna", "Día de Reparto", "Pendiente a entrega", "Total histórico", "Total adeudado"):
            self.tree.heading(col, text=col)
            self.tree.column(col, width=140, anchor="center")

        self.tree.pack(fill="both", expand=True, pady=16)

        self.label_total = self.crear_label_tema(frame, "", font=("Segoe UI", 12, "bold"))
        self.label_total.pack(pady=(8, 0))

        self.label_version = self.crear_label_tema(frame, "Versión 1.4.0", font=("Segoe UI", 10, "italic"), texto_color="secundario")
        self.label_version.pack(pady=(8, 0))

        self.ver_clientes()

    def registrar_widget_tema(self, widget, fondo="bg", texto="primario"):
        if not widget:
            return
        for idx, (w, _, _) in enumerate(self.widgets_tema):
            if w is widget:
                self.widgets_tema[idx] = (widget, fondo, texto)
                break
        else:
            self.widgets_tema.append((widget, fondo, texto))
        self._aplicar_colores_widget(widget, fondo, texto)

    def _aplicar_colores_widget(self, widget, fondo, texto):
        if not widget or not widget.winfo_exists():
            return
        colores = self.colores
        color_bg = None
        if fondo == "bg":
            color_bg = colores["bg"]
        elif fondo == "panel":
            color_bg = colores["panel"]
        elif fondo == "entry":
            color_bg = colores["entry_bg"]
        if color_bg is not None:
            try:
                widget.configure(bg=color_bg)
            except tk.TclError:
                pass
        if isinstance(widget, tk.Entry):
            widget.configure(bg=colores["entry_bg"], insertbackground=colores["texto"])
            if hasattr(widget, "_placeholder_text"):
                widget._placeholder_color = colores["placeholder"]
                widget._text_color = colores["entry_fg"]
                if widget.get() == widget._placeholder_text:
                    widget.configure(fg=widget._placeholder_color)
                else:
                    widget.configure(fg=widget._text_color)
            else:
                widget.configure(fg=colores["entry_fg"])
        if isinstance(widget, tk.Listbox):
            widget.configure(bg=colores["panel"], fg=colores["texto"], selectbackground=colores["tree_sel"], selectforeground=colores["tree_fg"])
        if isinstance(widget, tk.LabelFrame):
            widget.configure(bg=colores["panel"], fg=colores["texto"])
        if isinstance(widget, tk.Label):
            if texto == "primario":
                widget.configure(fg=colores["texto"])
            elif texto == "secundario":
                widget.configure(fg=colores["texto_secundario"])
        if isinstance(widget, tk.Toplevel):
            widget.configure(bg=colores["panel"])

    def actualizar_tema_widgets(self):
        for widget, fondo, texto in list(self.widgets_tema):
            if widget and widget.winfo_exists():
                self._aplicar_colores_widget(widget, fondo, texto)

    def configurar_estilos_ttk(self):
        colores = self.colores
        self.style.configure(
            "Treeview",
            background=colores["tree_bg"],
            fieldbackground=colores["tree_bg"],
            foreground=colores["tree_fg"],
            rowheight=28,
            bordercolor=colores["panel"],
            relief="flat"
        )
        self.style.map("Treeview", background=[("selected", colores["tree_sel"])], foreground=[("selected", colores["tree_fg"])])
        self.style.configure("Treeview.Heading", font=("Segoe UI", 11, "bold"), background=colores["tree_header_bg"], foreground=colores["tree_header_fg"])
        self.style.configure("TCombobox", fieldbackground=colores["entry_bg"], foreground=colores["entry_fg"], background=colores["entry_bg"])
        self.style.map("TCombobox", fieldbackground=[("readonly", colores["entry_bg"])], foreground=[("readonly", colores["entry_fg"])])
        self.style.configure("TButton", background=colores["boton_bg"], foreground=colores["boton_fg"], padding=5)
        self.style.map("TButton", background=[("active", colores["boton_hover"]), ("pressed", colores["boton_hover"])])
        self.style.configure("TLabel", background=colores["bg"], foreground=colores["texto"])
        self.style.configure(
            "Switch.TCheckbutton",
            background=colores["bg"],
            foreground=colores["texto"],
            padding=(4, 2),
            indicatoron=False
        )
        self.style.map(
            "Switch.TCheckbutton",
            background=[("selected", colores["boton_bg"]), ("active", colores["boton_hover"])],
            foreground=[("selected", colores["boton_fg"]), ("active", colores["boton_fg"])],
            relief=[("pressed", "sunken"), ("!pressed", "flat")]
        )

    def aplicar_tema_actual(self):
        self.colores = self.temas[self.tema_actual]
        self.configurar_estilos_ttk()
        self.root.configure(bg=self.colores["bg"])
        self.actualizar_tema_widgets()
        self.actualizar_texto_boton_tema()
        self.ver_clientes()

    def alternar_tema(self):
        seleccionado = False
        if hasattr(self, "tema_var"):
            seleccionado = bool(self.tema_var.get())
        nuevo_tema = "oscuro" if seleccionado else "claro"
        if nuevo_tema != self.tema_actual:
            self.tema_actual = nuevo_tema
            self.aplicar_tema_actual()
        else:
            if hasattr(self, "tema_var"):
                self.tema_var.set(self.tema_actual == "oscuro")

    def actualizar_texto_boton_tema(self):
        if hasattr(self, "boton_tema") and self.boton_tema:
            if hasattr(self, "tema_var"):
                self.tema_var.set(self.tema_actual == "oscuro")
            icono = "🌙" if self.tema_actual == "oscuro" else "☀️"
            self.boton_tema.config(text=icono)

    def crear_frame_tema(self, parent, fondo="bg", **kwargs):
        frame = tk.Frame(parent, **kwargs)
        self.registrar_widget_tema(frame, fondo=fondo, texto=None)
        return frame

    def crear_label_tema(self, parent, texto, font=None, fondo="bg", texto_color="primario", **kwargs):
        label = tk.Label(parent, text=texto, font=font, **kwargs)
        self.registrar_widget_tema(label, fondo=fondo, texto=texto_color)
        return label

    def crear_toplevel_tema(self, titulo, geometry=None, resizable=True):
        win = tk.Toplevel(self.root)
        win.title(titulo)
        if geometry:
            win.geometry(geometry)
        if not resizable:
            win.resizable(False, False)
        self.registrar_widget_tema(win, fondo="panel", texto=None)
        self.centrar_ventana(win)
        return win

    def registrar_descendencia_tema(self, widget):
        if not widget or not widget.winfo_exists():
            return
        for child in widget.winfo_children():
            if isinstance(child, tk.Entry):
                self.registrar_widget_tema(child, fondo="entry")
            elif isinstance(child, (tk.Frame, tk.LabelFrame)):
                self.registrar_widget_tema(child, fondo="panel", texto=None)
            elif isinstance(child, tk.Label):
                self.registrar_widget_tema(child, fondo="panel")
            elif isinstance(child, tk.Listbox):
                self.registrar_widget_tema(child, fondo="panel")
            elif isinstance(child, tk.Toplevel):
                self.registrar_widget_tema(child, fondo="panel", texto=None)
            self.registrar_descendencia_tema(child)

    # ------------------ Mostrar clientes ------------------

    def ver_clientes(self):
        for item in self.tree.get_children():
            self.tree.delete(item)

        self.actualizar_opciones_dias(self.filtro_dia_actual or "Todos")

        filtro_comuna = self.filtro_comuna_actual
        filtro_dia_norm = normalizar(self.filtro_dia_actual) if self.filtro_dia_actual else ""

        clientes_filtrados = []
        for cliente in self.data:
            if filtro_comuna and self.estandarizar_comuna(cliente.get("comuna")) != filtro_comuna:
                continue
            if filtro_dia_norm and normalizar(cliente.get("dia_reparto", "")) != filtro_dia_norm:
                continue
            clientes_filtrados.append(cliente)

        clientes_ordenados = sorted(clientes_filtrados, key=lambda c: c.get("cajas_de_huevos", 0), reverse=True)
        total_pendiente = 0

        for c in clientes_ordenados:
            pendiente = c.get("cajas_de_huevos", 0)
            total_pendiente += pendiente

            comuna_valor = self.estandarizar_comuna(c.get("comuna"))
            comuna_mostrar = comuna_valor or ""

            dia_valor = (c.get("dia_reparto", "") or "").strip()
            dia_mostrar = dia_valor.capitalize() if dia_valor else "-"

            precio = self.obtener_precio_por_comuna(comuna_valor)

            total_adeudado = pendiente * precio

            self.tree.insert("", "end", values=(
                c.get("nombre_completo", ""),
                c.get("telefono", ""),
                c.get("direccion", ""),
                comuna_mostrar,
                dia_mostrar,
                pendiente,
                c.get("cajas_de_huevos_total", 0),
                f"${total_adeudado:,.0f}".replace(",", ".")
            ))

        # Actualizar la etiqueta con el total de cajas pendientes
        filtros_activos = []
        if self.filtro_comuna_actual:
            filtros_activos.append(f"comuna {self.filtro_comuna_actual}")
        if self.filtro_dia_actual:
            filtros_activos.append(f"día {self.filtro_dia_actual}")
        resumen_filtros = f" • Filtros: {', '.join(filtros_activos)}" if filtros_activos else ""
        self.label_total.config(text=f"🥚 Total de cajas pendientes a entrega: {total_pendiente}{resumen_filtros}")
        self.label_total.pack(pady=(8, 0))

    def aplicar_placeholder(self, entry, placeholder_text):
        self.registrar_widget_tema(entry, fondo="entry")
        entry._placeholder_text = placeholder_text
        entry._placeholder_color = self.colores["placeholder"]
        entry._text_color = self.colores["entry_fg"]

        def on_focus_in(_):
            if entry.get() == entry._placeholder_text and entry.cget("fg") == entry._placeholder_color:
                entry.delete(0, tk.END)
                entry.config(fg=entry._text_color)

        def on_focus_out(_):
            if not entry.get():
                entry.insert(0, entry._placeholder_text)
                entry.config(fg=entry._placeholder_color)

        entry.bind("<FocusIn>", on_focus_in, add="+")
        entry.bind("<FocusOut>", on_focus_out, add="+")

        if not entry.get():
            entry.insert(0, entry._placeholder_text)
            entry.config(fg=entry._placeholder_color)
        elif entry.get() == entry._placeholder_text:
            entry.config(fg=entry._placeholder_color)
        else:
            entry.config(fg=entry._text_color)

    def obtener_valor_entry(self, entry):
        placeholder_text = getattr(entry, "_placeholder_text", None)
        placeholder_color = getattr(entry, "_placeholder_color", None)
        valor = entry.get().strip()
        if placeholder_text is not None and valor == placeholder_text and entry.cget("fg") == placeholder_color:
            return ""
        return valor

    def estandarizar_comuna(self, comuna):
        base = (comuna or "").strip()
        if not base:
            return ""
        candidato = base.title()
        clave = normalizar(candidato)
        existente = getattr(self, "_comunas_map", {}).get(clave)
        return existente or candidato

    def registrar_comuna(self, comuna, actualizar_opciones=True):
        nombre = self.estandarizar_comuna(comuna)
        if not nombre:
            return ""
        clave = normalizar(nombre)
        if clave not in self._comunas_map:
            self._comunas_map[clave] = nombre
            self.comunas = sorted(self._comunas_map.values())
            if actualizar_opciones:
                self.actualizar_opciones_comunas()
        return self._comunas_map[clave]

    def obtener_precio_por_comuna(self, comuna):
        clave = normalizar(self.estandarizar_comuna(comuna))
        if not clave:
            return PRECIO_CAJA
        for comuna_guardada, precio in self.precios_por_comuna.items():
            if normalizar(comuna_guardada) == clave:
                return precio
        return PRECIO_CAJA

    def actualizar_comunas_existentes(self, comunas_guardadas=None):
        self._comunas_map = {}

        if comunas_guardadas:
            for comuna in comunas_guardadas:
                self.registrar_comuna(comuna, actualizar_opciones=False)

        precios_ajustados = {}
        for comuna, precio in list(self.precios_por_comuna.items()):
            comuna_canonica = self.registrar_comuna(comuna, actualizar_opciones=False)
            if comuna_canonica:
                precios_ajustados[comuna_canonica] = precio
        self.precios_por_comuna = precios_ajustados

        for cliente in self.data:
            comuna_canonica = self.registrar_comuna(cliente.get("comuna"), actualizar_opciones=False)
            cliente["comuna"] = comuna_canonica

        self.comunas = sorted(self._comunas_map.values())
        self.actualizar_opciones_comunas()
        self.actualizar_opciones_dias()

    def actualizar_opciones_comunas(self, seleccion_preferida=None):
        opciones_base = list(self.comunas)
        combos_activos = []
        for combo in getattr(self, "_combobox_comunas", []):
            if not combo.winfo_exists():
                continue
            combos_activos.append(combo)
            opciones = list(opciones_base)
            if getattr(combo, "incluir_todas", False):
                opciones = ["Todas"] + opciones
            if getattr(combo, "permitir_agregar_comuna", True):
                opciones = opciones + ["Agregar comuna..."]
            combo["values"] = opciones

            valor_actual = combo.get()
            if valor_actual not in opciones:
                if getattr(combo, "incluir_todas", False) and "Todas" in opciones:
                    combo.set("Todas")
                elif opciones:
                    combo.set(opciones[0])

            if seleccion_preferida and seleccion_preferida in opciones and valor_actual == "Agregar comuna...":
                combo.set(seleccion_preferida)

        self._combobox_comunas = combos_activos

    def actualizar_opciones_dias(self, seleccion_preferida=None):
        if not hasattr(self, "combo_filtro_dia") or not self.combo_filtro_dia.winfo_exists():
            return
        dias_disponibles = sorted({
            (c.get("dia_reparto") or "").strip().title()
            for c in self.data
            if (c.get("dia_reparto") or "").strip()
        })
        opciones = ["Todos"] + dias_disponibles if dias_disponibles else ["Todos"]
        valor_prev = self.combo_filtro_dia.get()
        preferencia = seleccion_preferida or (valor_prev if valor_prev in opciones else None)
        self.combo_filtro_dia["values"] = opciones
        if preferencia and preferencia in opciones:
            self.combo_filtro_dia.set(preferencia)
            self.filtro_dia_actual = None if preferencia == "Todos" else preferencia
        else:
            self.combo_filtro_dia.set("Todos")
            self.filtro_dia_actual = None

    def restablecer_filtros(self, actualizar_tabla=True):
        self.filtro_comuna_actual = None
        self.filtro_dia_actual = None
        if hasattr(self, "combo_filtro_comuna") and self.combo_filtro_comuna.winfo_exists():
            if "Todas" in self.combo_filtro_comuna["values"]:
                self.combo_filtro_comuna.set("Todas")
        self.actualizar_opciones_dias("Todos")
        if actualizar_tabla:
            self.ver_clientes()

    def crear_combobox_comunas(self, parent, valor_inicial=None, permitir_agregar=True, incluir_todas=False, width=24):
        combo = ttk.Combobox(parent, state="readonly", width=width)
        combo.permitir_agregar_comuna = permitir_agregar
        combo.incluir_todas = incluir_todas
        self._combobox_comunas.append(combo)
        self.actualizar_opciones_comunas()

        if incluir_todas and "Todas" in combo["values"]:
            combo.set("Todas")

        if valor_inicial:
            valor_est = self.estandarizar_comuna(valor_inicial)
            if incluir_todas and valor_inicial == "Todas" and "Todas" in combo["values"]:
                combo.set("Todas")
            elif valor_est and valor_est in self.comunas:
                combo.set(valor_est)

        if permitir_agregar:
            def on_select(_):
                if combo.get() == "Agregar comuna...":
                    nueva = simpledialog.askstring("Nueva comuna", "Ingresa el nombre de la comuna:")
                    if not nueva:
                        self.actualizar_opciones_comunas()
                        return
                    nueva_est = self.estandarizar_comuna(nueva)
                    if not nueva_est:
                        messagebox.showerror("Error", "El nombre de la comuna no es válido.")
                        self.actualizar_opciones_comunas()
                        return
                    comuna_registrada = self.registrar_comuna(nueva_est)
                    combo.set(comuna_registrada)
                    self.actualizar_opciones_comunas(comuna_registrada)
                    self.guardar_estado()

            combo.bind("<<ComboboxSelected>>", on_select, add="+")

        return combo

    def obtener_comuna_combo(self, combo):
        valor = (combo.get() or "").strip()
        if not valor or valor == "Agregar comuna..." or valor == "Todas":
            return ""
        return self.registrar_comuna(valor, actualizar_opciones=False)

    def dialogo_seleccion_comuna(self, titulo, mensaje, incluir_todas=False):
        dialogo = self.crear_toplevel_tema(titulo, resizable=False)
        dialogo.grab_set()

        self.crear_label_tema(dialogo, mensaje, font=("Segoe UI", 11), fondo="panel").pack(padx=18, pady=(14, 8))
        combo = self.crear_combobox_comunas(dialogo, incluir_todas=incluir_todas, width=28)
        combo.pack(padx=18, pady=(0, 12))

        resultado = {"comuna": None}

        def aceptar():
            seleccion = combo.get()
            if not seleccion or seleccion == "Agregar comuna...":
                messagebox.showwarning("Selecciona una comuna", "Debes elegir una comuna válida.")
                return
            if incluir_todas and seleccion == "Todas":
                resultado["comuna"] = None
            else:
                resultado["comuna"] = self.registrar_comuna(seleccion, actualizar_opciones=False)
            dialogo.destroy()

        def cancelar():
            resultado["comuna"] = None
            dialogo.destroy()

        botones = self.crear_frame_tema(dialogo, fondo="panel")
        botones.pack(pady=(0, 12))
        ttk.Button(botones, text="Aceptar", command=aceptar).grid(row=0, column=0, padx=6)
        ttk.Button(botones, text="Cancelar", command=cancelar).grid(row=0, column=1, padx=6)

        self.registrar_descendencia_tema(dialogo)

        dialogo.wait_window()
        return resultado["comuna"]

    def guardar_estado(self):
        guardar_datos({
            "clientes": self.data,
            "precio_caja": PRECIO_CAJA,
            "precios_por_comuna": self.precios_por_comuna,
            "movimientos": self.movimientos,
            "caja_manual": self.caja_manual,
            "comunas": self.comunas
        })

    def ventana_caja(self):
        win = self.crear_toplevel_tema("Gestión de Caja", geometry="600x600")

        tk.Label(win, text="💰 Gestión de caja", bg="#f7f9fb", font=("Segoe UI", 14, "bold")).pack(pady=(10, 8))

        def formato_moneda(valor):
            try:
                return f"${float(valor):,.0f}".replace(",", ".")
            except (TypeError, ValueError):
                return "$0"

        resumen_label = tk.Label(win, text="", bg="#f7f9fb", font=("Segoe UI", 11))
        resumen_label.pack(pady=(0, 4))
        saldo_label = tk.Label(win, text="", bg="#f7f9fb", font=("Segoe UI", 11, "bold"))
        saldo_label.pack(pady=(0, 10))

        def calcular_totales():
            ingresos = egresos = otros = 0.0
            for mov in self.movimientos:
                tipo = (mov.get("tipo") or "").strip().lower()
                try:
                    monto = float(mov.get("monto", 0) or 0)
                except (TypeError, ValueError):
                    monto = 0.0
                if tipo == "ingreso":
                    ingresos += monto
                elif tipo == "egreso":
                    egresos += monto
                else:
                    otros += monto
            return ingresos, egresos, otros

        def actualizar_resumen():
            ingresos, egresos, otros = calcular_totales()
            resumen_label.config(
                text=(
                    f"Ingresos: {formato_moneda(ingresos)}   •   "
                    f"Egresos: {formato_moneda(egresos)}   •   "
                    f"Otros: {formato_moneda(otros)}"
                )
            )
            saldo_neto = ingresos - egresos
            saldo_label.config(text=f"Saldo neto (ingresos - egresos): {formato_moneda(saldo_neto)}")

        ttk.Separator(win, orient="horizontal").pack(fill="x", padx=14, pady=(4, 10))
        tk.Label(win, text="Registros guardados", bg="#f7f9fb", font=("Segoe UI", 12, "bold")).pack(anchor="w", padx=18)

        tree_frame = tk.Frame(win, bg="#f7f9fb")
        tree_frame.pack(fill="both", expand=True, padx=12, pady=(6, 4))
        columnas = ("Fecha", "Tipo", "Monto", "Descripción", "Referencia")
        tree = ttk.Treeview(tree_frame, columns=columnas, show="headings", height=8)
        for col in columnas:
            tree.heading(col, text=col)
            ancho = 120 if col in ("Fecha", "Tipo", "Monto") else 220
            tree.column(col, width=ancho, anchor="center" if col != "Descripción" else "w")
        tree.column("Descripción", anchor="w")
        tree.column("Referencia", anchor="w")

        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=scrollbar.set)
        tree.grid(row=0, column=0, sticky="nsew")
        scrollbar.grid(row=0, column=1, sticky="ns")
        tree_frame.columnconfigure(0, weight=1)
        tree_frame.rowconfigure(0, weight=1)

        label_resumen_registros = tk.Label(win, text="Registros guardados: 0", bg="#f7f9fb", font=("Segoe UI", 10))
        label_resumen_registros.pack(anchor="w", padx=18, pady=(0, 6))

        id_to_index = {}

        def refrescar_registros():
            id_to_index.clear()
            tree.delete(*tree.get_children())
            if not self.movimientos:
                label_resumen_registros.config(text="Registros guardados: 0")
                actualizar_resumen()
                return

            movimientos_ordenados = sorted(
                enumerate(self.movimientos),
                key=lambda par: par[1].get("fecha_iso", par[1].get("fecha", ""))
            )

            se_actualizo = False
            for idx_original, mov in movimientos_ordenados:
                if not mov.get("id"):
                    mov["id"] = datetime.now().strftime("%Y%m%d%H%M%S%f")
                    self.movimientos[idx_original] = mov
                    se_actualizo = True

                mov_id = mov.get("id")
                id_to_index[mov_id] = idx_original
                monto_valor = mov.get("monto", 0)
                descripcion = mov.get("descripcion", "")
                referencia = mov.get("metodo") or mov.get("cliente") or mov.get("referencia") or ""

                tree.insert(
                    "",
                    "end",
                    iid=mov_id,
                    values=(
                        mov.get("fecha", ""),
                        mov.get("tipo", ""),
                        formato_moneda(monto_valor),
                        descripcion,
                        referencia
                    )
                )

            label_resumen_registros.config(text=f"Registros guardados: {len(self.movimientos)}")
            if se_actualizo:
                self.guardar_estado()
            actualizar_resumen()

        def agregar_registro():
            self.abrir_formulario_registro(refrescar_registros)

        def eliminar_registro():
            sel = tree.selection()
            if not sel:
                messagebox.showerror("Error", "Selecciona un registro para eliminar.")
                return
            item_id = sel[0]
            if not messagebox.askyesno("Confirmar", "¿Eliminar el registro seleccionado?"):
                return
            idx_original = id_to_index.get(item_id)
            if idx_original is None or idx_original >= len(self.movimientos):
                messagebox.showerror("Error", "No se encontró el registro seleccionado.")
                return
            del self.movimientos[idx_original]
            self.guardar_estado()
            refrescar_registros()
            messagebox.showinfo("Éxito", "Registro eliminado correctamente.")

        registros_btn_frame = tk.Frame(win, bg="#f7f9fb")
        registros_btn_frame.pack(pady=(0, 10))
        ttk.Button(registros_btn_frame, text="Agregar registro", command=agregar_registro).grid(row=0, column=0, padx=6)
        ttk.Button(registros_btn_frame, text="Eliminar seleccionado", command=eliminar_registro).grid(row=0, column=1, padx=6)
        ttk.Button(registros_btn_frame, text="Cerrar", command=win.destroy).grid(row=0, column=2, padx=6)

        self.registrar_descendencia_tema(win)

        refrescar_registros()

    def abrir_formulario_registro(self, callback_refresco):
        win = self.crear_toplevel_tema("Nuevo registro de caja", geometry="360x360")

        tk.Label(win, text="Registrar movimiento manual", bg="#f7f9fb", font=("Segoe UI", 13, "bold")).pack(pady=(10, 10))

        tk.Label(win, text="Tipo de registro:", bg="#f7f9fb").pack(anchor="w", padx=18)
        tipos = ["Ingreso", "Egreso", "Otro"]
        combo_tipo = ttk.Combobox(win, state="readonly", values=tipos, width=18)
        combo_tipo.pack(pady=(0, 8))
        combo_tipo.current(0)

        tk.Label(win, text="Fecha (dd-mm-aaaa hh:mm):", bg="#f7f9fb").pack(anchor="w", padx=18)
        entry_fecha = tk.Entry(win, width=24, font=("Segoe UI", 10))
        fecha_actual = datetime.now().strftime("%d-%m-%Y %H:%M")
        entry_fecha.insert(0, fecha_actual)
        entry_fecha.pack(pady=(0, 8))

        tk.Label(win, text="Monto:", bg="#f7f9fb").pack(anchor="w", padx=18)
        entry_monto = tk.Entry(win, width=18, font=("Segoe UI", 10))
        entry_monto.pack(pady=(0, 8))
        self.aplicar_placeholder(entry_monto, "Ej: 25000")

        tk.Label(win, text="Descripción (opcional):", bg="#f7f9fb").pack(anchor="w", padx=18)
        entry_descripcion = tk.Entry(win, width=32, font=("Segoe UI", 10))
        entry_descripcion.pack(pady=(0, 8))
        self.aplicar_placeholder(entry_descripcion, "Ej: Venta en feria")

        tk.Label(win, text="Referencia / Método (opcional):", bg="#f7f9fb").pack(anchor="w", padx=18)
        entry_referencia = tk.Entry(win, width=32, font=("Segoe UI", 10))
        entry_referencia.pack(pady=(0, 12))
        self.aplicar_placeholder(entry_referencia, "Ej: Transferencia BancoEstado")

        def guardar_registro():
            tipo = combo_tipo.get() or "Otro"
            monto_texto = self.obtener_valor_entry(entry_monto)
            if not monto_texto:
                messagebox.showerror("Error", "Ingresa un monto para el registro.")
                return
            monto_limpio = monto_texto.replace(" ", "").replace("$", "").replace(".", "").replace(",", ".")
            try:
                monto = float(monto_limpio)
            except ValueError:
                messagebox.showerror("Error", "El monto debe ser numérico.")
                return
            if monto <= 0:
                messagebox.showerror("Error", "El monto debe ser mayor a 0.")
                return

            fecha_texto = self.obtener_valor_entry(entry_fecha) or fecha_actual
            try:
                fecha_iso = datetime.strptime(fecha_texto, "%d-%m-%Y %H:%M").isoformat()
            except ValueError:
                messagebox.showwarning("Advertencia", "Formato de fecha inválido. Se usará la fecha actual.")
                fecha_texto = datetime.now().strftime("%d-%m-%Y %H:%M")
                fecha_iso = datetime.now().isoformat()
            descripcion = self.obtener_valor_entry(entry_descripcion)
            referencia = self.obtener_valor_entry(entry_referencia)

            registro = {
                "id": datetime.now().strftime("%Y%m%d%H%M%S%f"),
                "fecha": fecha_texto,
                "fecha_iso": fecha_iso,
                "tipo": tipo,
                "monto": round(monto, 2),
                "descripcion": descripcion,
                "referencia": referencia
            }

            self.movimientos.append(registro)
            self.guardar_estado()
            callback_refresco()
            messagebox.showinfo("Éxito", "Registro guardado correctamente.")
            win.destroy()

        botones = tk.Frame(win, bg="#f7f9fb")
        botones.pack(pady=(6, 10))
        ttk.Button(botones, text="Guardar", command=guardar_registro).grid(row=0, column=0, padx=6)
        ttk.Button(botones, text="Cancelar", command=win.destroy).grid(row=0, column=1, padx=6)

        self.registrar_descendencia_tema(win)

        win.grab_set()
    # ------------------ Cambiar precio de la caja ------------------

    def cambiar_precio_caja(self):
        def guardar_precio():
            try:
                valor_ingresado = self.obtener_valor_entry(entry_precio)
                if not valor_ingresado:
                    messagebox.showerror("Error", "Ingrese un número válido.")
                    return
                nuevo_precio = int(valor_ingresado)
                if nuevo_precio <= 0:
                    messagebox.showerror("Error", "El precio debe ser mayor a 0.")
                    return
                global PRECIO_CAJA
                PRECIO_CAJA = nuevo_precio
                # Guardar el nuevo precio en los datos
                self.guardar_estado()
                self.ver_clientes()  # 🔹 Actualizar la tabla con el nuevo precio
                win.destroy()
                messagebox.showinfo("Éxito", f"El precio de la bandeja se actualizó a ${PRECIO_CAJA}.")
            except ValueError:
                messagebox.showerror("Error", "Ingrese un número válido.")

        win = self.crear_toplevel_tema("Cambiar Precio de la Bandeja", geometry="300x150")

        tk.Label(win, text="Precio actual: $" + str(PRECIO_CAJA), bg="#f7f9fb", font=("Segoe UI", 11)).pack(pady=(10, 4))
        tk.Label(win, text="Nuevo precio:", bg="#f7f9fb", font=("Segoe UI", 10)).pack(pady=(4, 0))

        entry_precio = tk.Entry(win, width=20, font=("Segoe UI", 10))
        entry_precio.pack(pady=6)
        self.aplicar_placeholder(entry_precio, "Ej: 7000")

        ttk.Button(win, text="Guardar", command=guardar_precio).pack(pady=10)

        self.registrar_descendencia_tema(win)

    # ------------------ Agregar cliente (placeholders) ------------------

    def ventana_agregar_cliente(self):
        def guardar_cliente():
            nombre = self.obtener_valor_entry(entry_nombre)
            telefono = self.obtener_valor_entry(entry_telefono)
            direccion = self.obtener_valor_entry(entry_direccion)
            comuna = self.obtener_comuna_combo(combo_comuna)
            dia_reparto = self.obtener_valor_entry(entry_dia_reparto)  # Nuevo campo para día de reparto

            if not nombre or not telefono or not direccion or not comuna:
                messagebox.showerror("Error", "Todos los campos excepto el día de reparto son obligatorios.")
                return

            nuevo_cliente = {
                "nombre_completo": nombre,
                "telefono": telefono,
                "direccion": direccion,
                "comuna": comuna,
                "cajas_de_huevos_total": 0,
                "cajas_de_huevos": 0,
                "dia_reparto": dia_reparto or None  # Guardar como None si está vacío
            }

            self.data.append(nuevo_cliente)
            self.actualizar_comunas_existentes(self.comunas)
            self.guardar_estado()
            self.ver_clientes()
            win.destroy()
            messagebox.showinfo("Éxito", "Cliente agregado correctamente.")

        win = self.crear_toplevel_tema("Agregar Cliente", geometry="400x400")

        tk.Label(win, text="Agregar Cliente", bg="#f7f9fb", font=("Segoe UI", 14, "bold")).pack(pady=(10, 10))

        tk.Label(win, text="Nombre Completo:", bg="#f7f9fb", font=("Segoe UI", 10)).pack(anchor="w", padx=20)
        entry_nombre = tk.Entry(win, width=30, font=("Segoe UI", 10))
        entry_nombre.pack(pady=5)
        self.aplicar_placeholder(entry_nombre, "Ej: Juan Pérez")

        tk.Label(win, text="Teléfono:", bg="#f7f9fb", font=("Segoe UI", 10)).pack(anchor="w", padx=20)
        entry_telefono = tk.Entry(win, width=30, font=("Segoe UI", 10))
        entry_telefono.pack(pady=5)
        self.aplicar_placeholder(entry_telefono, "Ej: 912345678")

        tk.Label(win, text="Dirección:", bg="#f7f9fb", font=("Segoe UI", 10)).pack(anchor="w", padx=20)
        entry_direccion = tk.Entry(win, width=30, font=("Segoe UI", 10))
        entry_direccion.pack(pady=5)
        self.aplicar_placeholder(entry_direccion, "Ej: Av. Siempre Viva 742")

        tk.Label(win, text="Comuna:", bg="#f7f9fb", font=("Segoe UI", 10)).pack(anchor="w", padx=20)
        combo_comuna = self.crear_combobox_comunas(win, width=28)
        combo_comuna.pack(pady=5)

        tk.Label(win, text="Día de Reparto (opcional):", bg="#f7f9fb", font=("Segoe UI", 10)).pack(anchor="w", padx=20)  # Etiqueta para el nuevo campo
        entry_dia_reparto = tk.Entry(win, width=30, font=("Segoe UI", 10))
        entry_dia_reparto.pack(pady=5)  # Campo de entrada para el día de reparto
        self.aplicar_placeholder(entry_dia_reparto, "Ej: Lunes")

        ttk.Button(win, text="Guardar", command=guardar_cliente).pack(pady=20)

        self.registrar_descendencia_tema(win)

    # ------------------ Nuevo pedido (placeholders + coincidencias) ------------------

    def ventana_nuevo_pedido(self):
        win = self.crear_toplevel_tema("Nuevo Pedido", geometry="500x450", resizable=False)

        tk.Label(win, text="Buscar cliente (nombre o telefono):", bg="#f7f9fb").pack(pady=(10, 0))
        entry_buscar = tk.Entry(win, width=44, font=("Segoe UI", 10))
        entry_buscar.pack(pady=(10, 5))
        self.aplicar_placeholder(entry_buscar, "Ej: María González o 912345678")

        tk.Label(win, text="Cantidad de cajas a agregar:", bg="#f7f9fb").pack(pady=(10, 0))
        entry_cantidad = tk.Entry(win, width=18, font=("Segoe UI", 10))
        entry_cantidad.pack(pady=6)
        self.aplicar_placeholder(entry_cantidad, "Ej: 10")

        # Crear listbox para mostrar resultados de búsqueda
        listbox = tk.Listbox(win, height=10, width=50)
        listbox.pack(pady=(5, 10))  # Ajustar margen superior e inferior

        resultados = []  # Lista para almacenar los resultados de búsqueda

        # Actualizar resultados en tiempo real mientras se escribe
        def actualizar_resultados(event):
            query = self.obtener_valor_entry(entry_buscar).lower()
            listbox.delete(0, tk.END)
            resultados.clear()
            for cliente in self.data:
                nombre = cliente.get("nombre_completo", "")
                comuna = (cliente.get("comuna", "") or "").strip()
                if query in nombre.lower() or query in cliente.get("telefono", ""):
                    resultados.append(cliente)
                    listbox.insert(tk.END, f"{nombre} - {comuna.capitalize() if comuna else ''}")
            if not resultados:
                listbox.insert(tk.END, "No se encontraron resultados.")

        # Asociar la actualización de resultados al evento de escritura
        entry_buscar.bind("<KeyRelease>", actualizar_resultados)

        # Cargar lista inicial de clientes
        actualizar_resultados(None)

        def agregar_pedido():
            sel = listbox.curselection()
            if not sel:
                messagebox.showerror("Error", "Seleccione un cliente.")
                return
            if not resultados or sel[0] >= len(resultados):
                messagebox.showerror("Error", "Seleccione un cliente válido de la lista.")
                return

            cliente = resultados[sel[0]]

            cantidad_texto = self.obtener_valor_entry(entry_cantidad)
            if not cantidad_texto:
                messagebox.showerror("Error", "El campo de cantidad no puede estar vacío.")
                return
            if not cantidad_texto.isdigit():
                messagebox.showerror("Error", "Ingrese un número válido para la cantidad de cajas.")
                return

            cantidad = int(cantidad_texto)
            if cantidad <= 0:
                messagebox.showerror("Error", "Ingrese una cantidad mayor que 0.")
                return
            cliente["cajas_de_huevos"] = cliente.get("cajas_de_huevos", 0) + cantidad
            cliente["cajas_de_huevos_total"] = cliente.get("cajas_de_huevos_total", 0) + cantidad
            self.guardar_estado()
            self.ver_clientes()
            messagebox.showinfo("Éxito", f"Se agregaron {cantidad} cajas a {cliente.get('nombre_completo','')}.")
            win.destroy()

        ttk.Button(win, text="Agregar pedido", command=agregar_pedido).pack(pady=10)

        self.registrar_descendencia_tema(win)

    # ------------------ Editar (búsqueda + opciones claras) ------------------

    def ventana_editar(self):
        win = self.crear_toplevel_tema("Editar Cliente o Pedido", geometry="460x420")

        tk.Label(win, text="Buscar cliente (nombre o telefono):", bg="#f7f9fb").pack(pady=(10, 0))
        entry_buscar = tk.Entry(win, width=46, font=("Segoe UI", 10))
        entry_buscar.pack(pady=6)
        self.aplicar_placeholder(entry_buscar, "Ej: Juan Pérez o 912345678")

        # Crear listbox para mostrar resultados de búsqueda
        listbox = tk.Listbox(win, height=10, width=70)  # 🔹 Definido correctamente
        listbox.pack(pady=6)
        resultados = []  # 🔹 Definido correctamente

        def buscar(event=None):
            listbox.delete(0, tk.END)
            texto = normalizar(self.obtener_valor_entry(entry_buscar))
            resultados.clear()
            if not texto:
                return
            for c in self.data:
                if texto in normalizar(c.get("nombre_completo", "")) or texto in normalizar(c.get("telefono", "")):
                    resultados.append(c)
                    listbox.insert(tk.END, f"{c.get('nombre_completo','')} ({c.get('telefono','')}) - {c.get('comuna','')} — Pendiente: {c.get('cajas_de_huevos',0)}")

        entry_buscar.bind("<KeyRelease>", buscar)

        def seleccionar():
            sel = listbox.curselection()
            if not sel:
                messagebox.showerror("Error", "Seleccione un cliente.")
                return
            cliente = resultados[sel[0]]

            # Ventana de opciones clara (Editar datos, Editar pedido, Eliminar)
            win_op = self.crear_toplevel_tema("¿Qué deseas hacer?", geometry="360x160", resizable=False)

            tk.Label(win_op, text=f"{cliente.get('nombre_completo','')}", bg="#f7f9fb", font=("Segoe UI", 11, "bold")).pack(pady=(10, 6))
            tk.Label(win_op, text=f"Pendiente: {cliente.get('cajas_de_huevos',0)} cajas — Comuna: {cliente.get('comuna','')}", bg="#f7f9fb").pack(pady=(0,8))

            def editar_datos():
                win_op.destroy()
                win.destroy()
                self.editar_datos_cliente(cliente)

            def editar_pedido():
                win_op.destroy()
                win.destroy()
                self.editar_pedido_cliente(cliente)

            def eliminar_cliente():
                if messagebox.askyesno("Confirmar eliminación", f"¿Eliminar a {cliente.get('nombre_completo','')}? Esta acción no se puede deshacer."):
                    try:
                        self.data.remove(cliente)
                        self.guardar_estado()
                        self.ver_clientes()
                        win_op.destroy()
                        win.destroy()
                        messagebox.showinfo("Eliminado", "Cliente eliminado correctamente.")
                    except Exception as e:
                        messagebox.showerror("Error", str(e))

            btn_frame = tk.Frame(win_op, bg="#f7f9fb")
            btn_frame.pack(pady=6)
            ttk.Button(btn_frame, text="Editar datos", width=12, command=editar_datos).grid(row=0, column=0, padx=6, pady=6)
            ttk.Button(btn_frame, text="Editar pedido", width=12, command=editar_pedido).grid(row=0, column=1, padx=6, pady=6)
            ttk.Button(btn_frame, text="Eliminar", width=12, command=eliminar_cliente).grid(row=1, column=0, columnspan=2, pady=6)
            ttk.Button(win_op, text="Cancelar", command=win_op.destroy).pack(pady=(4,6))

            self.registrar_descendencia_tema(win_op)

        ttk.Button(win, text="Seleccionar", command=seleccionar).pack(pady=8)

        self.registrar_descendencia_tema(win)

    # ------------------ Subventanas de edición ------------------

    def editar_datos_cliente(self, cliente):
        win = self.crear_toplevel_tema(f"Editar datos: {cliente.get('nombre_completo','')}", geometry="360x340")

        campos = [
            ("Nombre completo", "nombre_completo", cliente.get("nombre_completo", "")),
            ("telefono (opcional)", "telefono", cliente.get("telefono", "")),
            ("Dirección", "direccion", cliente.get("direccion", "")),
        ]

        entries = {}
        for label_text, key, valor in campos:
            tk.Label(win, text=label_text, bg="#f7f9fb").pack(pady=(8, 0))
            e = tk.Entry(win, width=42)
            e.insert(0, valor)
            e.pack()
            entries[key] = e

        tk.Label(win, text="Comuna", bg="#f7f9fb").pack(pady=(8, 0))
        combo_comuna = self.crear_combobox_comunas(win, valor_inicial=cliente.get("comuna"), width=28)
        combo_comuna.pack()

        def guardar():
            nombre = entries["nombre_completo"].get().strip()
            direccion = entries["direccion"].get().strip()
            comuna = self.obtener_comuna_combo(combo_comuna)
            if not nombre or not direccion or not comuna:
                messagebox.showerror("Error", "Complete los campos obligatorios.")
                return
            cliente["nombre_completo"] = nombre
            cliente["telefono"] = entries["telefono"].get().strip()
            cliente["direccion"] = direccion
            cliente["comuna"] = comuna
            self.actualizar_comunas_existentes(self.comunas)
            self.guardar_estado()
            self.ver_clientes()
            win.destroy()
            messagebox.showinfo("Éxito", "Datos del cliente actualizados.")

        ttk.Button(win, text="Guardar cambios", command=guardar).pack(pady=12)

        self.registrar_descendencia_tema(win)

    def editar_pedido_cliente(self, cliente):
        win = self.crear_toplevel_tema(f"Editar pedido: {cliente.get('nombre_completo','')}", geometry="340x240")

        tk.Label(
            win,
            text=f"Cliente: {cliente.get('nombre_completo', '')}",
            bg="#f7f9fb",
            font=("Segoe UI", 11, "bold")
        ).pack(pady=(12, 4))
        tk.Label(
            win,
            text=f"Pendiente actual: {cliente.get('cajas_de_huevos', 0)} cajas",
            bg="#f7f9fb",
            font=("Segoe UI", 10)
        ).pack(pady=(0, 10))

        def reemplazar():
            def guardar_reemplazo():
                try:
                    texto_nuevo = self.obtener_valor_entry(entry_reemplazar)
                    if not texto_nuevo:
                        messagebox.showerror("Error", "Ingrese una cantidad válida.")
                        return
                    nuevo = int(texto_nuevo)
                    if nuevo < 0:
                        messagebox.showerror("Error", "La cantidad no puede ser negativa.")
                        return
                    cliente["cajas_de_huevos"] = nuevo
                    # Ajustar histórico si es menor que total actual
                    cliente["cajas_de_huevos_total"] = max(cliente.get("cajas_de_huevos_total", 0), nuevo)
                    self.guardar_estado()
                    self.ver_clientes()
                    win_replace.destroy()
                    win.destroy()
                    messagebox.showinfo("Éxito", "Pedido reemplazado correctamente.")
                except ValueError:
                    messagebox.showerror("Error", "Ingrese un número válido.")
            win_replace = self.crear_toplevel_tema("Reemplazar cantidad pendiente", geometry="300x140", resizable=False)
            tk.Label(win_replace, text="Nueva cantidad pendiente:", bg="#f7f9fb").pack(pady=(10,4))
            entry_reemplazar = tk.Entry(win_replace, width=12)
            entry_reemplazar.insert(0, str(cliente.get("cajas_de_huevos",0)))
            entry_reemplazar.pack(pady=6)
            self.aplicar_placeholder(entry_reemplazar, "Ej: 5")
            ttk.Button(win_replace, text="Guardar", command=guardar_reemplazo).pack(pady=6)
            self.registrar_descendencia_tema(win_replace)

        btn_frame = tk.Frame(win, bg="#f7f9fb")
        btn_frame.pack(pady=8)
        ttk.Button(btn_frame, text="Editar cantidad", width=18, command=reemplazar).grid(row=0, column=0, padx=6, pady=6)
        ttk.Button(win, text="Cancelar", command=win.destroy).pack(pady=(6,8))

        self.registrar_descendencia_tema(win)

    # ------------------ Resumen / Estadísticas ------------------

    def ventana_resumen(self):
        if not self.data:
            messagebox.showinfo("Sin datos", "No hay clientes registrados.")
            return

        win = self.crear_toplevel_tema("📊 Resumen de Pedidos", geometry="540x540")

        total_pendiente = sum(c.get("cajas_de_huevos", 0) for c in self.data)
        clientes_con_pedido = [c for c in self.data if c.get("cajas_de_huevos", 0) > 0]
        total_clientes_pendientes = len(clientes_con_pedido)

        # Agrupar por comuna
        resumen_comunas = {}
        for c in clientes_con_pedido:
            comuna = c.get("comuna", "Sin comuna")
            resumen_comunas[comuna] = resumen_comunas.get(comuna, 0) + c.get("cajas_de_huevos", 0)

        resumen_ordenado = sorted(resumen_comunas.items(), key=lambda x: x[1], reverse=True)

        tk.Label(win, text=f"🥚 Total de cajas pendientes: {total_pendiente}", bg="#f7f9fb", font=("Segoe UI", 12, "bold")).pack(pady=(12, 6))
        tk.Label(win, text=f"👥 Clientes con pedidos activos: {total_clientes_pendientes}", bg="#f7f9fb", font=("Segoe UI", 10)).pack(pady=(0, 8))

        # Mostrar porcentaje por comuna
        if total_pendiente > 0:
            tk.Label(win, text="Porcentaje por comuna:", bg="#f7f9fb", font=("Segoe UI", 10, "underline")).pack(pady=(6,4))
            for comuna, total in resumen_ordenado:
                pct = (total / total_pendiente) * 100
                tk.Label(win, text=f"{comuna}: {total} cajas — {pct:.1f}%", bg="#f7f9fb").pack(anchor="w", padx=20)

        # Tabla con detalle por comuna
        tk.Label(win, text="Detalle por comuna:", bg="#f7f9fb", font=("Segoe UI", 10, "underline")).pack(pady=(10,4))
        tree = ttk.Treeview(win, columns=("Comuna", "Cajas pendientes", "Clientes con pedido"), show="headings", height=10)
        tree.heading("Comuna", text="Comuna")
        tree.heading("Cajas pendientes", text="Cajas pendientes")
        tree.heading("Clientes con pedido", text="Clientes con pedido")
        tree.column("Comuna", anchor="center", width=200)
        tree.column("Cajas pendientes", anchor="center", width=120)
        tree.column("Clientes con pedido", anchor="center", width=140)

        # calcular clientes por comuna
        clientes_por_comuna = {}
        for c in clientes_con_pedido:
            clientes_por_comuna[c.get("comuna", "Sin comuna")] = clientes_por_comuna.get(c.get("comuna", "Sin comuna"), 0) + 1

        for comuna, total in resumen_ordenado:
            tree.insert("", "end", values=(comuna, total, clientes_por_comuna.get(comuna, 0)))

        tree.pack(pady=8)

        ttk.Button(win, text="Cerrar", command=win.destroy).pack(pady=10)

        self.registrar_descendencia_tema(win)

    # ------------------ Generar reparto ------------------

    def generar_reparto(self):
        if not self.data:
            messagebox.showwarning("Sin clientes", "No hay clientes registrados.")
            return

        filtrar = messagebox.askyesno("Filtro", "¿Desea filtrar el reparto por comuna?")
        clientes_filtrados = self.data
        comuna = None

        if filtrar:
            if not self.comunas:
                messagebox.showwarning("Sin comunas", "No hay comunas registradas. Agrega una antes de filtrar.")
                return
            comuna_seleccionada = self.dialogo_seleccion_comuna("Filtrar por comuna", "Selecciona la comuna del reparto:")
            if not comuna_seleccionada:
                messagebox.showwarning("Advertencia", "Debe seleccionar una comuna válida.")
                return
            comuna = comuna_seleccionada
            clientes_filtrados = [c for c in self.data if self.estandarizar_comuna(c.get("comuna")) == comuna]
            if not clientes_filtrados:
                messagebox.showinfo("Sin resultados", f"No hay clientes en la comuna '{comuna}'.")
                return

        # Modificar la lógica de generación de reparto para incluir filtro por día
        filtrar_dia = messagebox.askyesno("Filtro", "¿Desea filtrar el reparto por día de reparto?")
        if filtrar_dia:
            dia_reparto = simpledialog.askstring("Filtrar por día", "Ingrese el día de reparto (por ejemplo, Lunes):")
            if not dia_reparto:
                messagebox.showwarning("Advertencia", "Debe ingresar un día válido.")
                return
            clientes_filtrados = [c for c in clientes_filtrados if normalizar(c.get("dia_reparto", "")) == normalizar(dia_reparto)]
            if not clientes_filtrados:
                messagebox.showinfo("Sin resultados", f"No hay clientes con día de reparto '{dia_reparto}'.")
                return

        clientes_con_pedidos = [c for c in clientes_filtrados if c.get("cajas_de_huevos", 0) > 0]
        if not clientes_con_pedidos:
            messagebox.showinfo("Sin pedidos", "No hay pedidos pendientes para generar reparto.")
            return

        fecha_actual = datetime.now().strftime("%d-%m-%Y")
        nombre_archivo = f"reparto_huevos_{(comuna or 'general').replace(' ', '_').lower()}_{fecha_actual}.xlsx"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reparto Huevos"

        # Definir encabezados y asegurarse de que `ws` esté correctamente utilizado
        encabezados = ["Nombre completo", "Teléfono", "Dirección", "Comuna", "Cajas de huevos", "Monto a pagar", "Pagado SI/NO", "Metodo de pago"]
        ws.append(encabezados)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        total_cajas = 0  # 🔹 Contador de total de cajas
        total_ganancias = 0  # 🔹 Contador de total a ganar

        for cliente in clientes_con_pedidos:
            cajas = cliente.get("cajas_de_huevos", 0)
            comuna_cliente = self.estandarizar_comuna(cliente.get("comuna"))
            precio = self.obtener_precio_por_comuna(comuna_cliente)
            monto_a_pagar = cajas * precio
            total_cajas += cajas
            total_ganancias += monto_a_pagar

            ws.append([
                cliente.get("nombre_completo", ""),
                cliente.get("telefono", ""),
                cliente.get("direccion", ""),
                comuna_cliente,
                cajas,
                f"${monto_a_pagar:,.0f}".replace(",", ".")  # Formato $000.000.000
            ])

        # 🔹 Nueva fila al final con el total
        ws.append([])
        ws.append(["", "", "", "Total", total_cajas, f"${total_ganancias:,.0f}".replace(",", ".")])

        # Ajustar ancho de columnas
        for column in ws.columns:
            max_length = max(len(str(cell.value)) if cell.value else 0 for cell in column)
            ws.column_dimensions[column[0].column_letter].width = max_length + 2

        # Guardar el archivo Excel
        wb.save(nombre_archivo)

        # Confirmar si marcar como entregados
        if messagebox.askyesno("Marcar entregidos", "¿Deseas marcar los pedidos generados como entregidos (poner 0)?"):
            for cliente in clientes_con_pedidos:
                cliente["cajas_de_huevos"] = 0
            # Preservar los precios por comuna al guardar los datos
            self.guardar_estado()
            self.ver_clientes()

        messagebox.showinfo("Éxito", f"Archivo '{nombre_archivo}' generado correctamente con total de {total_cajas} cajas y ganancias de ${total_ganancias:,.0f}".replace(",", "."))

    # ------------------ Centrar ventana ------------------

    def centrar_ventana(self, win):
        win.update_idletasks()
        width = win.winfo_width()
        height = win.winfo_height()
        x = (win.winfo_screenwidth() // 2) - (width // 2)
        y = (win.winfo_screenheight() // 2) - (height // 2)
        win.geometry(f"+{x}+{y}")

    # Crear una nueva ventana para gestionar precios por comuna
    def gestionar_precios_por_comuna(self):
        win = self.crear_toplevel_tema("Gestionar Precios por Comuna", geometry="460x520")

        def formatear_moneda(valor):
            try:
                return f"${float(valor):,.0f}".replace(",", ".")
            except (TypeError, ValueError):
                return "$0"

        tk.Label(win, text="Precios por Comuna", bg="#f7f9fb", font=("Segoe UI", 13, "bold")).pack(pady=(10, 4))
        tk.Label(
            win,
            text=f"Precio general actual: {formatear_moneda(PRECIO_CAJA)}",
            bg="#f7f9fb",
            font=("Segoe UI", 10)
        ).pack(pady=(0, 8))

        selector_frame = tk.LabelFrame(win, text="Asignar precio personalizado", bg="#f7f9fb", padx=8, pady=8)
        selector_frame.pack(fill="x", padx=12, pady=(0, 12))

        tk.Label(selector_frame, text="Comuna:", bg="#f7f9fb", font=("Segoe UI", 10)).grid(row=0, column=0, sticky="w")
        combo_comuna = self.crear_combobox_comunas(selector_frame, width=26)
        combo_comuna.grid(row=0, column=1, sticky="ew", padx=(8, 0))

        tk.Label(selector_frame, text="Precio personalizado:", bg="#f7f9fb", font=("Segoe UI", 10)).grid(row=1, column=0, sticky="w", pady=(6, 0))
        entry_precio = tk.Entry(selector_frame, width=18, font=("Segoe UI", 10))
        entry_precio.grid(row=1, column=1, sticky="w", padx=(8, 0), pady=(6, 0))
        self.aplicar_placeholder(entry_precio, "Ej: 7500")

        botones_selector = tk.Frame(selector_frame, bg="#f7f9fb")
        botones_selector.grid(row=2, column=0, columnspan=2, sticky="w", pady=(10, 0))
        btn_guardar_selector = ttk.Button(botones_selector, text="Guardar precio", width=18)
        btn_guardar_selector.pack(side="left", padx=(0, 8))
        btn_restablecer_selector = ttk.Button(botones_selector, text="Usar precio general", width=20)
        btn_restablecer_selector.pack(side="left")

        selector_frame.columnconfigure(1, weight=1)

        tree_frame = tk.Frame(win, bg="#f7f9fb")
        tree_frame.pack(fill="both", expand=True, padx=12, pady=(0, 8))

        columnas = ("Comuna", "Precio", "Origen")
        tree = ttk.Treeview(tree_frame, columns=columnas, show="headings", height=12)
        tree.heading("Comuna", text="Comuna")
        tree.heading("Precio", text="Precio actual")
        tree.heading("Origen", text="Origen del precio")
        tree.column("Comuna", anchor="center", width=180)
        tree.column("Precio", anchor="center", width=130)
        tree.column("Origen", anchor="center", width=130)

        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=scrollbar.set)
        tree.grid(row=0, column=0, sticky="nsew")
        scrollbar.grid(row=0, column=1, sticky="ns")
        tree_frame.columnconfigure(0, weight=1)
        tree_frame.rowconfigure(0, weight=1)

        resumen_label = tk.Label(win, text="", bg="#f7f9fb", font=("Segoe UI", 10))
        resumen_label.pack(anchor="w", padx=14)

        def actualizar_entry_para_comuna(comuna_objetivo):
            if not comuna_objetivo:
                placeholder_text = getattr(entry_precio, "_placeholder_text", "")
                placeholder_color = getattr(entry_precio, "_placeholder_color", "#777")
                entry_precio.delete(0, tk.END)
                if placeholder_text:
                    entry_precio.insert(0, placeholder_text)
                entry_precio.config(fg=placeholder_color)
                return
            if comuna_objetivo in self.precios_por_comuna:
                entry_precio.delete(0, tk.END)
                entry_precio.insert(0, str(self.precios_por_comuna[comuna_objetivo]))
                entry_precio.config(fg=getattr(entry_precio, "_text_color", "#000"))
            else:
                placeholder_text = getattr(entry_precio, "_placeholder_text", "")
                placeholder_color = getattr(entry_precio, "_placeholder_color", "#777")
                entry_precio.delete(0, tk.END)
                if placeholder_text:
                    entry_precio.insert(0, placeholder_text)
                entry_precio.config(fg=placeholder_color)

        def refrescar_tree(seleccionar_actual=None):
            tree.delete(*tree.get_children())
            comunas_union = set(self.comunas)
            comunas_union.update(self.estandarizar_comuna(c) for c in self.precios_por_comuna.keys())
            comunas_ordenadas = sorted(c for c in comunas_union if c)
            personalizados = 0
            for comuna in comunas_ordenadas:
                if comuna in self.precios_por_comuna:
                    precio = self.precios_por_comuna[comuna]
                    origen = "Personalizado"
                    personalizados += 1
                else:
                    precio = PRECIO_CAJA
                    origen = "General"
                tree.insert("", "end", iid=comuna, values=(comuna, formatear_moneda(precio), origen))
            resumen_label.config(
                text=f"Comunas registradas: {len(comunas_ordenadas)} • Personalizados: {personalizados}"
            )
            if seleccionar_actual and tree.exists(seleccionar_actual):
                tree.selection_set(seleccionar_actual)
                tree.see(seleccionar_actual)

        refrescar_tree()

        def guardar_precio_personalizado():
            comuna_sel = self.obtener_comuna_combo(combo_comuna)
            if not comuna_sel:
                messagebox.showerror("Error", "Selecciona una comuna para asignar un precio.")
                return
            valor_ingresado = self.obtener_valor_entry(entry_precio)
            if not valor_ingresado:
                messagebox.showerror("Error", "Ingresa un precio para la comuna seleccionada.")
                return
            digitos = ''.join(ch for ch in valor_ingresado if ch.isdigit())
            if not digitos:
                messagebox.showerror("Error", "El precio debe contener solo números.")
                return
            precio = int(digitos)
            if precio <= 0:
                messagebox.showerror("Error", "El precio debe ser mayor a 0.")
                return
            comuna_registrada = self.registrar_comuna(comuna_sel)
            self.precios_por_comuna[comuna_registrada] = precio
            self.guardar_estado()
            self.ver_clientes()
            combo_comuna.set(comuna_registrada)
            refrescar_tree(seleccionar_actual=comuna_registrada)
            actualizar_entry_para_comuna(comuna_registrada)
            messagebox.showinfo(
                "Éxito",
                f"Precio personalizado para {comuna_registrada} guardado en {formatear_moneda(precio)}."
            )

        def restablecer_precio_general():
            comuna_sel = self.obtener_comuna_combo(combo_comuna)
            if not comuna_sel:
                messagebox.showerror("Error", "Selecciona una comuna para restablecer el precio general.")
                return
            comuna_registrada = self.registrar_comuna(comuna_sel, actualizar_opciones=False)
            if comuna_registrada not in self.precios_por_comuna:
                messagebox.showinfo("Sin cambios", "La comuna ya usa el precio general.")
                actualizar_entry_para_comuna(comuna_registrada)
                return
            if not messagebox.askyesno(
                "Confirmar",
                f"¿Restablecer el precio general para {comuna_registrada}?"
            ):
                return
            self.precios_por_comuna.pop(comuna_registrada, None)
            self.guardar_estado()
            self.ver_clientes()
            refrescar_tree(seleccionar_actual=comuna_registrada)
            actualizar_entry_para_comuna(comuna_registrada)
            messagebox.showinfo(
                "Éxito",
                f"{comuna_registrada} volverá a usar el precio general ({formatear_moneda(PRECIO_CAJA)})."
            )

        def on_tree_select(_):
            seleccion = tree.selection()
            if not seleccion:
                return
            comuna_sel = tree.item(seleccion[0], "values")[0]
            combo_comuna.set(comuna_sel)
            actualizar_entry_para_comuna(comuna_sel)

        tree.bind("<<TreeviewSelect>>", on_tree_select)

        def on_combobox_change(_):
            comuna_sel = self.obtener_comuna_combo(combo_comuna)
            refrescar_tree(seleccionar_actual=comuna_sel)
            actualizar_entry_para_comuna(comuna_sel)

        combo_comuna.bind("<<ComboboxSelected>>", on_combobox_change, add="+")

        entry_precio.bind("<Return>", lambda _: guardar_precio_personalizado())
        btn_guardar_selector.config(command=guardar_precio_personalizado)
        btn_restablecer_selector.config(command=restablecer_precio_general)

        botones = tk.Frame(win, bg="#f7f9fb")
        botones.pack(pady=12)
        ttk.Button(botones, text="Guardar precio", command=guardar_precio_personalizado).grid(row=0, column=0, padx=6)
        ttk.Button(botones, text="Restablecer general", command=restablecer_precio_general).grid(row=0, column=1, padx=6)
        ttk.Button(botones, text="Cerrar", command=win.destroy).grid(row=0, column=2, padx=6)

        actualizar_entry_para_comuna(self.obtener_comuna_combo(combo_comuna))

        self.registrar_descendencia_tema(win)

    # ------------------ Agregar día de reparto ------------------

    def agregar_dia_reparto(self):
        def guardar_dia():
            dia = self.obtener_valor_entry(entry_dia)
            if dia:
                cliente["dia_reparto"] = dia
                self.guardar_estado()
                self.ver_clientes()
                win.destroy()
                messagebox.showinfo("Éxito", f"Día de reparto para {cliente['nombre_completo']} actualizado a '{dia}'.")
            else:
                messagebox.showwarning("Advertencia", "No se ingresó ningún día de reparto.")

        sel = self.tree.selection()
        if not sel:
            messagebox.showerror("Error", "Seleccione un cliente para agregar el día de reparto.")
            return

        cliente = self.data[self.tree.index(sel[0])]

        win = self.crear_toplevel_tema("Agregar Día de Reparto", geometry="300x150")

        tk.Label(win, text=f"Cliente: {cliente['nombre_completo']}", bg="#f7f9fb", font=("Segoe UI", 11)).pack(pady=(10, 4))
        tk.Label(win, text="Día de reparto:", bg="#f7f9fb", font=("Segoe UI", 10)).pack(pady=(4, 0))

        entry_dia = tk.Entry(win, width=20, font=("Segoe UI", 10))
        valor_actual = cliente.get("dia_reparto", "")
        if valor_actual:
            entry_dia.insert(0, valor_actual)
        entry_dia.pack(pady=6)
        self.aplicar_placeholder(entry_dia, "Ej: Lunes")

        ttk.Button(win, text="Guardar", command=guardar_dia).pack(pady=10)

        self.registrar_descendencia_tema(win)

# ------------------ Ejecución ------------------

if __name__ == "__main__":
    root = tk.Tk()
    app = App(root)
    root.mainloop()  # 🔹 Verificado: formato correcto
